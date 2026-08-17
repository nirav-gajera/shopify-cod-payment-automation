import os
import sys
import glob
import argparse
import requests
import openpyxl
from dotenv import load_dotenv

load_dotenv()

SHOP = os.getenv("SHOPIFY_SHOP")
CLIENT_ID = os.getenv("SHOPIFY_CLIENT_ID")
CLIENT_SECRET = os.getenv("SHOPIFY_CLIENT_SECRET")

GRAPHQL_URL = f"https://{SHOP}.myshopify.com/admin/api/2026-04/graphql.json"
ORDERS_DIR = os.path.join(os.path.dirname(__file__), "orders")


# ── Auth ──────────────────────────────────────────────────────────────────────

def get_access_token():
    token_url = f"https://{SHOP}.myshopify.com/admin/oauth/access_token"
    resp = requests.post(
        token_url,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        data={
            "grant_type": "client_credentials",
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
        },
    )
    resp.raise_for_status()
    return resp.json()["access_token"]


# ── Read order numbers from all xlsx files ────────────────────────────────────

def collect_order_numbers() -> list:
    order_numbers = []
    xlsx_files = glob.glob(os.path.join(ORDERS_DIR, "*.xlsx"))
    print(f"Found {len(xlsx_files)} Excel file(s) in orders/")

    for path in xlsx_files:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active

        header = None
        order_id_col = None

        for row in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip() if c else "" for c in row]
                for idx, h in enumerate(header):
                    if h.lower() == "orderid":
                        order_id_col = idx
                        break
                if order_id_col is None:
                    print(f"  [SKIP] No 'OrderId' column in {os.path.basename(path)}")
                    break
                continue

            val = row[order_id_col]
            if val is not None and str(val).strip():
                order_numbers.append(str(val).strip())

        wb.close()

    # Deduplicate while preserving order
    seen = set()
    unique = []
    for n in order_numbers:
        if n not in seen:
            seen.add(n)
            unique.append(n)

    print(f"Total unique order numbers collected: {len(unique)}")
    return unique


# ── GraphQL helpers ───────────────────────────────────────────────────────────

FIND_ORDER_QUERY = """
query FindOrder($query: String!) {
    orders(first: 1, query: $query) {
        nodes {
            id
            name
            displayFinancialStatus
        }
    }
}
"""

MARK_AS_PAID_MUTATION = """
mutation OrderMarkAsPaid($input: OrderMarkAsPaidInput!) {
    orderMarkAsPaid(input: $input) {
        order {
            id
            name
            displayFinancialStatus
        }
        userErrors {
            field
            message
        }
    }
}
"""


def gql(token, query, variables):
    resp = requests.post(
        GRAPHQL_URL,
        headers={
            "Content-Type": "application/json",
            "X-Shopify-Access-Token": token,
        },
        json={"query": query, "variables": variables},
    )
    resp.raise_for_status()
    return resp.json()


def find_order_id(token, order_number):
    """Returns (gid, financial_status) or (None, None) if not found."""
    data = gql(token, FIND_ORDER_QUERY, {"query": f"name:#{order_number}"})
    nodes = data.get("data", {}).get("orders", {}).get("nodes", [])
    if not nodes:
        return None, None
    node = nodes[0]
    return node["id"], node.get("displayFinancialStatus")


def mark_as_paid(token, order_gid):
    return gql(token, MARK_AS_PAID_MUTATION, {"input": {"id": order_gid}})


# ── Main ──────────────────────────────────────────────────────────────────────

def process_order(token, order_number):
    """Find and mark a single order as paid. Returns result status string."""
    print(f"Processing order #{order_number} ...", end=" ")

    gid, status = find_order_id(token, order_number)

    if gid is None:
        print("NOT FOUND")
        return "not_found"

    print(f"\n  GID    : {gid}")
    print(f"  Status : {status}")

    if status == "PAID":
        print("  -> already PAID, skipping.")
        return "already_paid"

    resp = mark_as_paid(token, gid)
    errors = resp.get("data", {}).get("orderMarkAsPaid", {}).get("userErrors", [])

    if errors:
        msgs = "; ".join(e["message"] for e in errors)
        print(f"  -> ERROR: {msgs}")
        return "error"

    new_status = (
        resp.get("data", {})
        .get("orderMarkAsPaid", {})
        .get("order", {})
        .get("displayFinancialStatus", "?")
    )
    print(f"  -> Marked as paid! New status: {new_status}")
    return "success"


def process_file(token, path):
    """Read all order numbers from one xlsx and mark them as paid."""
    print(f"\nFile: {os.path.basename(path)}")
    print("-" * 50)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header = None
    order_id_col = None
    order_numbers = []

    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip() if c else "" for c in row]
            for idx, h in enumerate(header):
                if h.lower() == "orderid":
                    order_id_col = idx
                    break
            if order_id_col is None:
                print("  [SKIP] No 'OrderId' column found.")
                break
            continue
        val = row[order_id_col]
        if val is not None and str(val).strip():
            order_numbers.append(str(val).strip())

    wb.close()

    results = {"success": 0, "already_paid": 0, "not_found": 0, "error": 0}
    for num in order_numbers:
        r = process_order(token, num)
        results[r] = results.get(r, 0) + 1

    print(f"\n  Summary -> Paid: {results['success']}  Already paid: {results['already_paid']}  "
          f"Not found: {results['not_found']}  Errors: {results['error']}")
    return results


def main():
    parser = argparse.ArgumentParser(
        description="Mark Shopify orders as paid from Excel invoice sheets."
    )
    group = parser.add_mutually_exclusive_group()
    group.add_argument(
        "--file", "-f",
        metavar="PATH",
        help="Path to a specific .xlsx file (relative or absolute)",
    )
    group.add_argument(
        "--order", "-o",
        metavar="ORDER_NUMBER",
        help="Single order number to mark as paid (e.g. 6689)",
    )
    # No args = process all files in orders/
    args = parser.parse_args()

    token = get_access_token()
    print("Access token obtained.")

    if args.order:
        # ── Single order number mode ──
        result = process_order(token, args.order)
        print(f"\nResult: {result}")

    elif args.file:
        # ── Single file mode ──
        path = args.file if os.path.isabs(args.file) else os.path.join(os.path.dirname(__file__), args.file)
        if not os.path.exists(path):
            print(f"ERROR: File not found: {path}")
            sys.exit(1)
        process_file(token, path)

    else:
        # ── All files mode ──
        xlsx_files = sorted(glob.glob(os.path.join(ORDERS_DIR, "*.xlsx")))
        print(f"Found {len(xlsx_files)} file(s) in orders/\n")
        totals = {"success": 0, "already_paid": 0, "not_found": 0, "error": 0}
        for path in xlsx_files:
            r = process_file(token, path)
            for k in totals:
                totals[k] += r.get(k, 0)
        print("\n" + "=" * 50)
        print("TOTAL SUMMARY")
        print("=" * 50)
        print(f"  Paid         : {totals['success']}")
        print(f"  Already paid : {totals['already_paid']}")
        print(f"  Not found    : {totals['not_found']}")
        print(f"  Errors       : {totals['error']}")


if __name__ == "__main__":
    main()
